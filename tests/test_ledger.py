"""Regression tests for money, records, imports, documents, and recovery."""
import base64
from datetime import date
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import zipfile

from app import Store, LedgerServer, business_cents, csv_report, money_cents, pdf_report, report_data, word_report


class LedgerTests(unittest.TestCase):
    def setUp(self):
        test_root=Path(__file__).resolve().parents[1]/".test-data"
        test_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=test_root)
        self.store = Store(Path(self.temp.name) / "ledger")

    def tearDown(self):
        self.temp.cleanup()

    def entry(self, **overrides):
        value = {"type": "income", "date": "2026-10-15", "name": "Example Brand", "source": "ShopMy",
                 "amount": "500.00", "fees": "20.00", "refunds": "10.00", "category": "Affiliate commission"}
        value.update(overrides)
        return self.store.save_transaction(value)["id"]

    def upload(self, name="Example receipt.txt", content=b"Synthetic sample receipt", **extra):
        return self.store.upload({"name": name, "base64": base64.b64encode(content).decode(), **extra})["id"]

    def test_exact_money_validation_and_rounding(self):
        self.assertEqual(money_cents("$1,234.56"), 123456)
        self.assertEqual(business_cents(101, 50), 51)
        for invalid in ("NaN", "Infinity", "-1", "0.001", "abc", None):
            with self.subTest(value=invalid), self.assertRaises(ValueError):
                money_cents(invalid)

    def test_gross_fees_refunds_business_use_and_quarter_boundary(self):
        self.entry()
        self.entry(type="expense", name="Example art vendor", amount="150", category="Artwork & content purchases",
                   business_percent=60, tax_review=True)
        self.entry(date="2026-09-30", amount="1000")
        report = report_data(self.store, "2026-10-01", "2026-12-31")
        self.assertEqual(len(report["entries"]), 2)
        self.assertEqual(report["totals"]["net_income"], 47000)
        self.assertEqual(report["totals"]["expense_total"], 15000)
        self.assertEqual(report["totals"]["business_expenses"], 9000)
        self.assertEqual(report["totals"]["result"], 38000)
        self.assertEqual(report["needs_attention"], 1)
        with self.assertRaises(ValueError):
            report_data(self.store, "2026-12-31", "2026-01-01")

    def test_invalid_entry_does_not_change_records(self):
        for changes in ({"fees": "600"}, {"date": "2026-02-30"}, {"business_percent": 101}, {"type": "other"}, {"name": " "}, {"amount": "0"}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.entry(**changes)
        self.assertEqual(self.store.snapshot()["transactions"], [])

    def test_original_document_dedup_links_edit_and_recoverable_removal(self):
        doc = self.upload()
        again = self.upload(name="renamed.txt")
        self.assertEqual(doc, again)
        self.assertEqual(len(self.store.snapshot()["documents"]), 1)
        ident = self.entry(type="expense", name="Sample vendor", amount="42.50", category="Content supplies", document_ids=[doc])
        entry = self.store.snapshot()["transactions"][0]
        self.assertEqual(entry["attention"], [])
        self.entry(id=ident, type="expense", name="Updated vendor", amount="43", document_ids=[doc], category="Content supplies")
        self.assertEqual(self.store.snapshot()["transactions"][0]["name"], "Updated vendor")
        self.store.remove_transaction(ident)
        self.assertEqual(self.store.snapshot()["transactions"], [])
        self.assertEqual(len(self.store.snapshot()["documents"]), 1)
        self.store.undo_remove(ident)
        self.assertEqual(len(self.store.snapshot()["transactions"]), 1)
        with self.store.connect() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM history").fetchone()[0], 2)

    def test_import_validation_is_atomic_and_duplicates_are_optional(self):
        valid = {"type":"income", "date":"2026-10-01", "name":"Example platform", "amount":"12.34"}
        with self.assertRaises(ValueError):
            self.store.csv_import({"entries": [valid, {**valid, "amount":"NaN"}]})
        self.assertEqual(len(self.store.snapshot()["transactions"]), 0)
        self.assertEqual(self.store.csv_import({"entries":[valid,valid]}), {"imported":1,"skipped":1})
        self.assertEqual(self.store.csv_import({"entries":[valid]}), {"imported":0,"skipped":1})
        self.assertEqual(self.store.csv_import({"entries":[valid],"skip_duplicates":False}), {"imported":1,"skipped":0})
        preview=self.store.csv_preview({"text":'\ufeffDate,Name,Amount\n2026-10-01,"Example, Inc.",12.34\n'})
        self.assertEqual(preview["rows"][0]["Name"], "Example, Inc.")

    def test_full_backup_restores_documents_settings_history_and_links(self):
        doc=self.upload(name="sample.pdf", content=b"%PDF-1.4 synthetic test", kind="Bank statement", tax_year=2026)
        ident=self.entry(document_ids=[doc])
        self.store.save_settings({"owner":"Example Owner"})
        backup=self.store.backup()
        self.entry(name="This change should be replaced")
        self.store.restore({"base64":base64.b64encode(backup).decode()})
        state=self.store.snapshot()
        self.assertEqual(len(state["transactions"]),1)
        self.assertEqual(state["transactions"][0]["id"],ident)
        self.assertEqual(state["settings"]["owner"],"Example Owner")
        self.assertEqual(state["documents"][0]["transaction_ids"],[ident])
        self.assertEqual(self.store.document(doc)[1].read_bytes(),b"%PDF-1.4 synthetic test")
        self.assertGreaterEqual(len(list(self.store.backups.glob("*.zip"))),2)

    def test_restore_rejects_zip_traversal_and_keeps_current_records(self):
        ident=self.entry()
        output=io.BytesIO()
        with zipfile.ZipFile(output,"w") as archive:
            archive.writestr("backup-info.json",json.dumps({"app":"Leah's Ledger","version":1}))
            archive.writestr("../escape.txt","malicious")
        with self.assertRaises(ValueError):
            self.store.restore({"base64":base64.b64encode(output.getvalue()).decode()})
        self.assertEqual(self.store.snapshot()["transactions"][0]["id"],ident)
        self.assertFalse((Path(self.temp.name)/"escape.txt").exists())

    def test_restore_rejects_tampered_document(self):
        self.upload()
        backup=io.BytesIO(self.store.backup())
        changed=io.BytesIO()
        with zipfile.ZipFile(backup) as source,zipfile.ZipFile(changed,"w") as dest:
            for name in source.namelist():
                dest.writestr(name,b"changed" if name.startswith("documents/") else source.read(name))
        with self.assertRaises(ValueError):
            self.store.restore({"base64":base64.b64encode(changed.getvalue()).decode()})
        self.assertEqual(len(self.store.snapshot()["documents"]),1)

    def test_year_statements_included_even_without_transaction_link(self):
        current=self.upload(name="current.txt",content=b"current",kind="Bank statement",tax_year=2026)
        self.upload(name="old.txt",content=b"old",kind="Bank statement",tax_year=2025)
        self.upload(name="unlinked receipt.txt",content=b"receipt",tax_year=2026)
        report=report_data(self.store,"2026-10-01","2026-12-31")
        self.assertEqual([d["id"] for d in report["documents"]],[current])

    def test_pdf_word_and_csv_exports_with_unicode_long_notes(self):
        self.entry(name="Example Café & Co.",notes="Review art business use. "*100,tax_review=True)
        self.entry(type="expense",name="=HYPERLINK(example)",amount="35.01",category="Content supplies",business_percent=50)
        report=report_data(self.store,"2026-01-01","2026-12-31")
        pdf=pdf_report(report)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(pdf))
        text="\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("Example Café & Co.",text)
        self.assertIn("$452.49",text)
        self.assertIn("Document index",text)
        from docx import Document
        doc=Document(io.BytesIO(word_report(report)))
        self.assertTrue(any("Example Café" in p.text for p in doc.paragraphs))
        self.assertIn("$452.49",str([[c.text for c in row.cells] for table in doc.tables for row in table.rows]))
        csv_text=csv_report(report).decode("utf-8-sig")
        self.assertIn("'=HYPERLINK(example)",csv_text)
        self.assertIn("Recorded business expense",csv_text)

    def test_http_authentication_origin_and_reports(self):
        self.entry()
        server=LedgerServer(("127.0.0.1",0),self.store)
        thread=threading.Thread(target=server.serve_forever,daemon=True)
        thread.start()
        try:
            base=f"http://127.0.0.1:{server.server_port}"
            with self.assertRaises(urllib.error.HTTPError) as denied:
                urllib.request.urlopen(base+"/api/state")
            self.assertEqual(denied.exception.code,403)
            for origin in (None,"https://unrelated.example"):
                headers={"X-Ledger-Token":server.token}
                if origin:headers["Origin"]=origin
                request=urllib.request.Request(base+"/api/state",headers=headers)
                if origin:
                    with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(request)
                else:
                    with urllib.request.urlopen(request) as response:
                        self.assertEqual(len(json.load(response)["transactions"]),1)
            request=urllib.request.Request(base+"/api/report?start=2026-01-01&end=2026-12-31&format=zip",headers={"X-Ledger-Token":server.token})
            with urllib.request.urlopen(request) as response:
                with zipfile.ZipFile(io.BytesIO(response.read())) as archive:
                    self.assertIn("Summary.pdf",archive.namelist())
                    self.assertIn("Summary.docx",archive.namelist())
            request=urllib.request.Request(base+"/api/backup",headers={"X-Ledger-Token":server.token})
            with urllib.request.urlopen(request) as response:
                backup=response.read()
            self.entry(name="Replace this later")
            request=urllib.request.Request(base+"/api/restore",data=backup,headers={"X-Ledger-Token":server.token,"Content-Type":"application/zip"})
            with urllib.request.urlopen(request) as response:
                self.assertTrue(json.load(response)["restored"])
            self.assertEqual(len(self.store.snapshot()["transactions"]),1)
        finally:
            server.shutdown();server.server_close();thread.join()


if __name__=="__main__":
    unittest.main()
