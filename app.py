"""Leah's Ledger: a private, offline Windows bookkeeping companion."""
from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import io
import json
import logging
import mimetypes
import os
from pathlib import Path
import secrets
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import zipfile
from datetime import date, datetime
from contextlib import contextmanager, closing
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

APP_NAME = "Leah's Ledger"
RESOURCES = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
SOURCES = ["Amazon Associates", "ShopMy", "Olivia Rose James · Royals", "Other"]
EXPENSE_CATEGORIES = ["Artwork & content purchases", "Content supplies", "Software & subscriptions",
                      "Equipment", "Advertising", "Professional services", "Office supplies",
                      "Phone & internet", "Travel", "Meals", "Other / ask my preparer"]
INCOME_CATEGORIES = ["Affiliate commission", "Sponsored content", "Ad revenue", "Other income"]
MAX_FILE = 25 * 1024 * 1024


def money_cents(value, *, allow_zero=True):
    try:
        number = Decimal(str(value).replace(",", "").replace("$", "").strip())
        if not number.is_finite() or number < 0 or number > Decimal("999999999.99"):
            raise ValueError()
        if number != number.quantize(Decimal("0.01")):
            raise ValueError()
        cents = int(number * 100)
        if not allow_zero and not cents:
            raise ValueError()
        return cents
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("Enter a valid amount with no more than two decimal places.") from None


def business_cents(amount, percent):
    return int((Decimal(amount) * Decimal(str(percent)) / 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def validate_date(value):
    try:
        return date.fromisoformat(str(value)).isoformat()
    except (ValueError, TypeError):
        raise ValueError("Choose a valid date.") from None


def fingerprint(entry):
    fields = [entry["date"], entry["type"], entry["name"].strip().casefold(), str(entry["amount_cents"])]
    return hashlib.sha256("|".join(fields).encode()).hexdigest()


def normalize_entry(data):
    kind = data.get("type")
    if kind not in ("income", "expense"):
        raise ValueError("Choose money in or money out.")
    entry = {"date": validate_date(data.get("date")), "type": kind}
    for key in ("name", "source", "category", "campaign", "notes", "account"):
        entry[key] = str(data.get(key, "")).strip()[:4000 if key == "notes" else 200]
    entry["payment_method"] = str(data.get("payment_method", "Direct deposit"))[:100]
    if not entry["name"]:
        raise ValueError("Enter who paid you or where you bought it.")
    entry["amount_cents"] = money_cents(data.get("amount"), allow_zero=False)
    entry["fees_cents"] = money_cents(data.get("fees", 0)) if kind == "income" else 0
    entry["refunds_cents"] = money_cents(data.get("refunds", 0)) if kind == "income" else 0
    if entry["fees_cents"] + entry["refunds_cents"] > entry["amount_cents"]:
        raise ValueError("Fees and refunds cannot exceed the gross income for this entry.")
    try:
        percent = Decimal(str(data.get("business_percent", 100)))
        if not percent.is_finite() or not 0 <= percent <= 100 or percent != percent.quantize(Decimal("0.01")):
            raise ValueError()
    except (InvalidOperation, ValueError):
        raise ValueError("Business use must be between 0 and 100%, with up to two decimal places.") from None
    entry["business_percent"] = float(percent) if kind == "expense" else 100
    entry["tax_review"] = int(bool(data.get("tax_review")))
    entry["receipt_not_needed"] = int(bool(data.get("receipt_not_needed")))
    entry["fingerprint"] = fingerprint(entry)
    return entry


class Store:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.documents = self.root / "documents"
        self.documents.mkdir(exist_ok=True)
        self.backups = self.root / "backups"
        self.backups.mkdir(exist_ok=True)
        self.db_path = self.root / "ledger.sqlite3"
        self.lock = threading.RLock()
        self._last_change_backup = None
        with self.connect() as con:
            con.executescript("""
                CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS transactions (
                    id TEXT PRIMARY KEY, date TEXT NOT NULL, type TEXT NOT NULL,
                    name TEXT NOT NULL, source TEXT NOT NULL, category TEXT NOT NULL,
                    campaign TEXT NOT NULL, notes TEXT NOT NULL, account TEXT NOT NULL,
                    payment_method TEXT NOT NULL, amount_cents INTEGER NOT NULL,
                    fees_cents INTEGER NOT NULL, refunds_cents INTEGER NOT NULL,
                    business_percent REAL NOT NULL, tax_review INTEGER NOT NULL,
                    receipt_not_needed INTEGER NOT NULL, fingerprint TEXT NOT NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, deleted INTEGER NOT NULL DEFAULT 0
                );
                CREATE INDEX IF NOT EXISTS transaction_date ON transactions(date);
                CREATE INDEX IF NOT EXISTS transaction_fingerprint ON transactions(fingerprint);
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, mime TEXT NOT NULL,
                    size INTEGER NOT NULL, sha256 TEXT NOT NULL UNIQUE, stored_name TEXT NOT NULL,
                    kind TEXT NOT NULL, created_at TEXT NOT NULL,
                    tax_year INTEGER NOT NULL DEFAULT 2026
                );
                CREATE TABLE IF NOT EXISTS transaction_documents (
                    transaction_id TEXT NOT NULL REFERENCES transactions(id),
                    document_id TEXT NOT NULL REFERENCES documents(id),
                    PRIMARY KEY (transaction_id, document_id)
                );
                CREATE TABLE IF NOT EXISTS history (
                    id INTEGER PRIMARY KEY, transaction_id TEXT NOT NULL,
                    action TEXT NOT NULL, snapshot TEXT NOT NULL, created_at TEXT NOT NULL
                );
            """)
            con.execute("INSERT OR IGNORE INTO settings VALUES ('owner', ?)", (json.dumps("Leah"),))
            columns = {r[1] for r in con.execute("PRAGMA table_info(documents)")}
            if "tax_year" not in columns:
                con.execute("ALTER TABLE documents ADD COLUMN tax_year INTEGER NOT NULL DEFAULT 2026")
        self.auto_backup()

    @contextmanager
    def connect(self):
        con = sqlite3.connect(self.db_path, timeout=20)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA foreign_keys=ON")
        try:
            with con:
                yield con
        finally:
            con.close()

    def settings(self):
        with self.connect() as con:
            return {r["key"]: json.loads(r["value"]) for r in con.execute("SELECT * FROM settings")}

    def save_settings(self, data):
        owner = str(data.get("owner", "Leah")).strip()[:80] or "Leah"
        with self.lock, self.connect() as con:
            con.execute("INSERT OR REPLACE INTO settings VALUES ('owner', ?)", (json.dumps(owner),))
        return {"owner": owner}

    def snapshot(self):
        with self.lock, self.connect() as con:
            entries = [dict(row) for row in con.execute("SELECT * FROM transactions WHERE deleted=0 ORDER BY date DESC, created_at DESC")]
            docs = [dict(row) for row in con.execute("SELECT * FROM documents ORDER BY created_at DESC")]
            links = [dict(row) for row in con.execute("SELECT td.* FROM transaction_documents td JOIN transactions t ON t.id=td.transaction_id WHERE t.deleted=0")]
        for entry in entries:
            entry["document_ids"] = [link["document_id"] for link in links if link["transaction_id"] == entry["id"]]
            entry["attention"] = []
            if not entry["category"]:
                entry["attention"].append("Choose a category")
            if entry["type"] == "expense" and not entry["document_ids"] and not entry["receipt_not_needed"]:
                entry["attention"].append("Add a receipt")
            if entry["tax_review"]:
                entry["attention"].append("Preparer review")
            if entry["type"] == "income" and not entry["source"]:
                entry["attention"].append("Choose an income source")
        for doc in docs:
            doc["transaction_ids"] = [link["transaction_id"] for link in links if link["document_id"] == doc["id"]]
            doc.pop("stored_name", None)
        return {"transactions": entries, "documents": docs, "settings": self.settings(),
                "sources": SOURCES, "expense_categories": EXPENSE_CATEGORIES,
                "income_categories": INCOME_CATEGORIES, "data_folder": str(self.root),
                "last_backup": self.last_backup()}

    def save_transaction(self, data):
        entry = normalize_entry(data)
        ident = str(data.get("id") or secrets.token_hex(12))
        document_ids = list(dict.fromkeys(data.get("document_ids", [])))
        if any(not isinstance(item, str) for item in document_ids):
            raise ValueError("Invalid document selection.")
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock, self.connect() as con:
            previous = con.execute("SELECT * FROM transactions WHERE id=? AND deleted=0", (ident,)).fetchone()
            if data.get("id") and previous is None:
                raise ValueError("This entry no longer exists. Refresh and try again.")
            for doc_id in document_ids:
                if con.execute("SELECT id FROM documents WHERE id=?", (doc_id,)).fetchone() is None:
                    raise ValueError("A selected document is no longer available.")
            if previous:
                con.execute("INSERT INTO history(transaction_id,action,snapshot,created_at) VALUES (?,?,?,?)",
                            (ident, "edit", json.dumps(dict(previous)), now))
                assignments = ",".join(f"{key}=?" for key in entry)
                con.execute(f"UPDATE transactions SET {assignments},updated_at=? WHERE id=?",
                            [*entry.values(), now, ident])
            else:
                full = {"id": ident, **entry, "created_at": now, "updated_at": now}
                con.execute(f"INSERT INTO transactions ({','.join(full)}) VALUES ({','.join('?' for _ in full)})", list(full.values()))
            con.execute("DELETE FROM transaction_documents WHERE transaction_id=?", (ident,))
            con.executemany("INSERT INTO transaction_documents VALUES (?,?)", [(ident, d) for d in document_ids])
        return {"id": ident}

    def remove_transaction(self, ident):
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock, self.connect() as con:
            previous = con.execute("SELECT * FROM transactions WHERE id=? AND deleted=0", (ident,)).fetchone()
            if previous is None:
                raise ValueError("This entry is already removed.")
            con.execute("INSERT INTO history(transaction_id,action,snapshot,created_at) VALUES (?,?,?,?)", (ident, "remove", json.dumps(dict(previous)), now))
            con.execute("UPDATE transactions SET deleted=1,updated_at=? WHERE id=?", (now, ident))
        return {"id": ident}

    def undo_remove(self, ident):
        with self.lock, self.connect() as con:
            con.execute("UPDATE transactions SET deleted=0,updated_at=? WHERE id=?", (datetime.now().isoformat(timespec="seconds"), ident))
        return {"id": ident}

    def upload(self, data):
        name = Path(str(data.get("name", "Document")).replace("\\", "/")).name[:200]
        if not name or name in (".", ".."):
            raise ValueError("Choose a file with a valid name.")
        try:
            contents = base64.b64decode(data.get("base64", ""), validate=True)
        except (ValueError, TypeError):
            raise ValueError("This file could not be read. Please choose it again.") from None
        if not contents or len(contents) > MAX_FILE:
            raise ValueError("Choose a non-empty file up to 25 MB.")
        digest = hashlib.sha256(contents).hexdigest()
        ident = secrets.token_hex(12)
        kind = str(data.get("kind", "Receipt"))[:80]
        try:
            tax_year = int(data.get("tax_year", date.today().year))
            if not 2000 <= tax_year <= 2100:
                raise ValueError()
        except (ValueError, TypeError):
            raise ValueError("Choose a valid year for this document.") from None
        suffix = Path(name).suffix.lower()
        stored = ident + suffix if len(suffix) < 15 else ident
        with self.lock, self.connect() as con:
            previous = con.execute("SELECT id FROM documents WHERE sha256=?", (digest,)).fetchone()
            if previous:
                return {"id": previous["id"], "duplicate": True}
            destination = self.documents / stored
            destination.write_bytes(contents)
            try:
                con.execute("INSERT INTO documents (id,name,mime,size,sha256,stored_name,kind,created_at,tax_year) VALUES (?,?,?,?,?,?,?,?,?)", (ident, name, mimetypes.guess_type(name)[0] or "application/octet-stream", len(contents), digest, stored, kind, datetime.now().isoformat(timespec="seconds"), tax_year))
            except Exception:
                destination.unlink(missing_ok=True)
                raise
        return {"id": ident, "duplicate": False}

    def document(self, ident):
        with self.connect() as con:
            row = con.execute("SELECT * FROM documents WHERE id=?", (ident,)).fetchone()
        if row is None:
            raise ValueError("Document not found.")
        path = self.documents / row["stored_name"]
        if not path.is_file():
            raise ValueError("The original document is missing. Restore it from a backup.")
        return dict(row), path

    def last_backup(self):
        files = sorted(self.backups.glob("*.zip"), key=lambda p: p.stat().st_mtime, reverse=True)
        return datetime.fromtimestamp(files[0].stat().st_mtime).isoformat(timespec="seconds") if files else None

    def backup(self, destination=None):
        with self.lock, tempfile.TemporaryDirectory(dir=self.root) as tmp:
            target = Path(tmp) / "ledger.sqlite3"
            with self.connect() as source, closing(sqlite3.connect(target)) as dest:
                source.backup(dest)
            output = io.BytesIO() if destination is None else destination
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(target, "ledger.sqlite3")
                for path in sorted(self.documents.iterdir()):
                    if path.is_file():
                        archive.write(path, "documents/" + path.name)
                archive.writestr("backup-info.json", json.dumps({"app": APP_NAME, "version": 1, "created": datetime.now().isoformat()}))
            return output.getvalue() if destination is None else None

    def auto_backup(self, force=False, on_change=False):
        today = date.today().isoformat()
        if not force:
            if on_change and self._last_change_backup == today:
                return
            if not on_change and list(self.backups.glob(f"{today}*.zip")):
                return
        destination = self.backups / f"{today}-{datetime.now():%H%M%S}-{secrets.token_hex(2)}.zip"
        temp = destination.with_suffix(".tmp")
        self.backup(temp)
        temp.replace(destination)
        if on_change:
            self._last_change_backup = today
        # Keep the most recent 30 daily/recovery backups. User exports are never removed.
        for old in sorted(self.backups.glob("*.zip"), key=lambda p: p.stat().st_mtime, reverse=True)[30:]:
            old.unlink()

    def restore(self, data):
        try:
            payload = base64.b64decode(data.get("base64", ""), validate=True)
        except (ValueError, TypeError):
            raise ValueError("Choose a valid Leah's Ledger backup ZIP.") from None
        if len(payload) > 250 * 1024 * 1024:
            raise ValueError("This backup is too large to restore here.")
        return self.restore_archive(io.BytesIO(payload))

    def restore_archive(self, source):
        with self.lock, tempfile.TemporaryDirectory(dir=self.root) as tmp:
            staging = Path(tmp)
            try:
                with zipfile.ZipFile(source) as archive:
                    if archive.getinfo("backup-info.json").file_size > 4096:
                        raise ValueError()
                    info = json.loads(archive.read("backup-info.json"))
                    if info.get("app") != APP_NAME or info.get("version") != 1:
                        raise ValueError()
                    if sum(i.file_size for i in archive.infolist()) > 4 * 1024 * 1024 * 1024:
                        raise ValueError()
                    for item in archive.infolist():
                        parts = item.filename.split("/")
                        allowed = item.filename in ("ledger.sqlite3", "backup-info.json") or (len(parts) == 2 and parts[0] == "documents" and parts[1] not in ("", ".", "..") and "\\" not in parts[1] and ":" not in parts[1])
                        if not allowed or item.is_dir():
                            raise ValueError()
                        destination = staging.joinpath(*parts)
                        destination.parent.mkdir(exist_ok=True)
                        with archive.open(item) as incoming, destination.open("wb") as outgoing:
                            shutil.copyfileobj(incoming, outgoing, 1024 * 1024)
                with closing(sqlite3.connect(staging / "ledger.sqlite3")) as restored:
                    if restored.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                        raise ValueError()
                    restored.row_factory = sqlite3.Row
                    for table in ("settings", "transactions", "documents", "transaction_documents", "history"):
                        restored.execute(f"SELECT * FROM {table} LIMIT 1")
                    for doc in restored.execute("SELECT * FROM documents"):
                        stored = doc["stored_name"]
                        if Path(stored).name != stored or "/" in stored or "\\" in stored or ":" in stored:
                            raise ValueError()
                        file = staging / "documents" / stored
                        if not file.is_file():
                            raise ValueError()
                        with file.open("rb") as contents:
                            if hashlib.file_digest(contents, "sha256").hexdigest() != doc["sha256"]:
                                raise ValueError()
            except (zipfile.BadZipFile, KeyError, json.JSONDecodeError, sqlite3.Error, ValueError, OSError):
                raise ValueError("This backup is incomplete or invalid. Your current records have not been changed.") from None
            self.auto_backup(force=True)
            (staging / "documents").mkdir(exist_ok=True)
            old_db, old_docs = staging / "old.sqlite3", staging / "old-documents"
            self.db_path.replace(old_db)
            try:
                self.documents.replace(old_docs)
                (staging / "ledger.sqlite3").replace(self.db_path)
                (staging / "documents").replace(self.documents)
            except Exception:
                if self.db_path.exists():
                    self.db_path.unlink()
                old_db.replace(self.db_path)
                if old_docs.exists():
                    if self.documents.exists():
                        shutil.rmtree(self.documents)
                    old_docs.replace(self.documents)
                raise
        return {"restored": True}

    def csv_preview(self, data):
        raw = str(data.get("text", "")).lstrip("\ufeff")
        if len(raw) > 5 * 1024 * 1024:
            raise ValueError("Choose a CSV file under 5 MB.")
        try:
            dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(raw), dialect=dialect)
        headers = reader.fieldnames or []
        if not headers or len(headers) != len(set(headers)):
            raise ValueError("The CSV needs a header row with unique column names.")
        rows = list(reader)
        if len(rows) > 5000:
            raise ValueError("Import up to 5,000 rows at a time.")
        if any(None in row for row in rows):
            raise ValueError("Some CSV rows have more values than the header. Check the file and try again.")
        return {"headers": headers, "rows": rows}

    def csv_import(self, data):
        rows = data.get("entries", [])
        if not isinstance(rows, list) or len(rows) > 5000:
            raise ValueError("Import up to 5,000 entries at a time.")
        normalized = [(normalize_entry(row), row) for row in rows]
        count = skipped = 0
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock, self.connect() as con:
            existing = {r[0] for r in con.execute("SELECT fingerprint FROM transactions WHERE deleted=0")}
            for entry, original in normalized:
                if data.get("skip_duplicates", True) and entry["fingerprint"] in existing:
                    skipped += 1
                    continue
                ident = secrets.token_hex(12)
                full = {"id": ident, **entry, "created_at": now, "updated_at": now}
                con.execute(f"INSERT INTO transactions ({','.join(full)}) VALUES ({','.join('?' for _ in full)})", list(full.values()))
                existing.add(entry["fingerprint"])
                count += 1
        return {"imported": count, "skipped": skipped}


def report_data(store, start, end):
    start, end = validate_date(start), validate_date(end)
    if start > end:
        raise ValueError("The report end date must be after the start date.")
    snapshot = store.snapshot()
    entries = sorted((e for e in snapshot["transactions"] if start <= e["date"] <= end), key=lambda e: (e["date"], e["created_at"]))
    totals = {"gross": 0, "fees": 0, "refunds": 0, "expense_total": 0, "business_expenses": 0}
    by_source, by_category = {}, {}
    for e in entries:
        if e["type"] == "income":
            totals["gross"] += e["amount_cents"]
            totals["fees"] += e["fees_cents"]
            totals["refunds"] += e["refunds_cents"]
            source = e["source"] or "Unassigned"
            by_source[source] = by_source.get(source, 0) + e["amount_cents"]
        else:
            totals["expense_total"] += e["amount_cents"]
            business = business_cents(e["amount_cents"], e["business_percent"])
            totals["business_expenses"] += business
            category = e["category"] or "Uncategorized"
            by_category[category] = by_category.get(category, 0) + business
    totals["net_income"] = totals["gross"] - totals["fees"] - totals["refunds"]
    totals["result"] = totals["net_income"] - totals["business_expenses"]
    ids = {e["id"] for e in entries}
    # Statements and tax forms apply to a year and may support many transactions.
    docs = [d for d in snapshot["documents"] if ids.intersection(d["transaction_ids"]) or
            (d["kind"] in ("Bank statement", "Tax form", "Platform statement") and
             int(start[:4]) <= d["tax_year"] <= int(end[:4]))]
    return {"start": start, "end": end, "owner": snapshot["settings"].get("owner", "Leah"), "entries": entries,
            "documents": docs, "totals": totals, "by_source": by_source, "by_category": by_category,
            "needs_attention": sum(bool(e["attention"]) for e in entries)}


def money(value):
    return f"{'-' if value < 0 else ''}${abs(value) / 100:,.2f}"


def csv_report(report):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["ID", "Date", "Type", "Payer or vendor", "Source", "Category", "Campaign", "Gross or purchase amount", "Fees", "Refunds", "Net deposit", "Business use %", "Recorded business expense", "Payment method", "Account label", "Review flags", "Documents", "Notes"])
    for e in report["entries"]:
        names = [d["name"] for d in report["documents"] if d["id"] in e["document_ids"]]
        row = [e["id"], e["date"], e["type"], e["name"], e["source"], e["category"], e["campaign"],
               f'{e["amount_cents"]/100:.2f}', f'{e["fees_cents"]/100:.2f}', f'{e["refunds_cents"]/100:.2f}',
               f'{(e["amount_cents"]-e["fees_cents"]-e["refunds_cents"])/100:.2f}' if e["type"] == "income" else "",
               e["business_percent"], f'{business_cents(e["amount_cents"],e["business_percent"])/100:.2f}' if e["type"] == "expense" else "",
               e["payment_method"], e["account"], "; ".join(e["attention"]), "; ".join(names), e["notes"]]
        # Prevent spreadsheet programs interpreting user text as formulas.
        writer.writerow(["'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v for v in row])
    return output.getvalue().encode("utf-8-sig")


def report_sections(report):
    t = report["totals"]
    summary = [("Gross income", t["gross"]), ("Fees withheld", t["fees"]), ("Refunds", t["refunds"]),
               ("Net deposits recorded", t["net_income"]), ("Total purchases", t["expense_total"]),
               ("Recorded business portion of purchases", t["business_expenses"]), ("Bookkeeping result before tax adjustments", t["result"])]
    return [("Period summary", [(k, money(v)) for k, v in summary]),
            ("Gross income by source", [(k, money(v)) for k, v in sorted(report["by_source"].items())]),
            ("Recorded business expenses by category", [(k, money(v)) for k, v in sorted(report["by_category"].items())])]


REPORT_NOTE = "Bookkeeping records for preparer review. Recorded business use is user supplied and does not establish deductibility. Artwork, equipment, and other flagged purchases may require different tax treatment. Fees withheld from income are included once in this report. Dates use the received/paid cash-recording approach; unpaid commissions are not included."


def detail_lines(entry, documents):
    e = entry
    lines = [f'{e["date"]} | {"Money in" if e["type"] == "income" else "Money out"} | {e["name"]}',
             f'Source: {e["source"] or "General business"} | Category: {e["category"] or "Uncategorized"}',
             f'Gross / purchase: {money(e["amount_cents"])}']
    if e["type"] == "income":
        lines.append(f'Fees: {money(e["fees_cents"])} | Refunds: {money(e["refunds_cents"])} | Net deposit: {money(e["amount_cents"]-e["fees_cents"]-e["refunds_cents"])}')
    else:
        lines.append(f'Business use: {e["business_percent"]:g}% | Recorded business portion: {money(business_cents(e["amount_cents"], e["business_percent"]))}')
    lines.append(f'Payment: {e["payment_method"] or "Unspecified"} | Account: {e["account"] or "Unspecified"}')
    if e["campaign"]:
        lines.append("Campaign: " + e["campaign"])
    if e["attention"]:
        lines.append("Needs attention: " + "; ".join(e["attention"]))
    if e["notes"]:
        lines.append("Notes: " + e["notes"])
    attached = [d["name"] for d in documents if d["id"] in e["document_ids"]]
    lines.append("Documents: " + ("; ".join(attached) if attached else "None linked"))
    lines.append("Record ID: " + e["id"])
    return lines


def pdf_report(report):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape
    font_folder = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    # Register under a lock: reportlab's font registry is shared by HTTP threads.
    if "Ledger" not in pdfmetrics.getRegisteredFontNames():
        regular, bold = font_folder / "segoeui.ttf", font_folder / "segoeuib.ttf"
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("Ledger", str(regular)))
            pdfmetrics.registerFont(TTFont("LedgerBold", str(bold)))
    font = "Ledger" if "Ledger" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
    bold = "LedgerBold" if font == "Ledger" else "Helvetica-Bold"
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=44, leftMargin=44, topMargin=54, bottomMargin=48,
                            title=f"{APP_NAME} | {report['start']} to {report['end']}", author=report["owner"])
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("LedgerTitle", fontName=bold, fontSize=24, leading=29, textColor=colors.HexColor("#45543e"), spaceAfter=10))
    styles.add(ParagraphStyle("LedgerHeading", fontName=bold, fontSize=13, leading=18, textColor=colors.HexColor("#526c4e"), spaceBefore=18, spaceAfter=8, keepWithNext=True))
    styles.add(ParagraphStyle("LedgerBody", fontName=font, fontSize=9, leading=14, spaceAfter=5, wordWrap="CJK"))
    styles.add(ParagraphStyle("LedgerDetail", fontName=font, fontSize=9, leading=13, spaceAfter=3, wordWrap="CJK"))
    styles.add(ParagraphStyle("LedgerEntry", fontName=bold, fontSize=11, leading=15, textColor=colors.HexColor("#526c4e"), spaceBefore=10, spaceAfter=7))
    styles.add(ParagraphStyle("LedgerSmall", fontName=font, fontSize=8, leading=12, textColor=colors.HexColor("#65705c"), spaceAfter=8, wordWrap="CJK"))
    p = lambda text, style="LedgerBody": Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])
    story = [p("Business income and expense report", "LedgerTitle"), p(f"{report['owner']} · {report['start']} through {report['end']}"),
             p(f"{len(report['entries'])} transactions · {report['needs_attention']} entries need attention · Generated {datetime.now():%Y-%m-%d %H:%M}", "LedgerSmall"), p(REPORT_NOTE, "LedgerSmall")]
    for title, rows in report_sections(report):
        story.append(p(title, "LedgerHeading"))
        rows = rows or [("No records in this period", "—")]
        table = Table([[p(k), p(v)] for k, v in rows], colWidths=[390, 134], hAlign="LEFT")
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f4e9")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LINEBELOW", (0, 0), (-1, -1), .4, colors.white)]))
        story.append(table)
    if not report["entries"]:
        story.append(p("Transaction details", "LedgerHeading"))
        story.append(p("No transactions in this period."))
    for index, e in enumerate(report["entries"]):
        lines = detail_lines(e, report["documents"])
        block = [p(lines[0], "LedgerEntry"), *[p(line, "LedgerDetail") for line in lines[1:]], Spacer(1, 8)]
        if index == 0:
            block.insert(0, p("Transaction details", "LedgerHeading"))
        story.append(KeepTogether(block))
    if report["documents"]:
        story.append(p("Document index", "LedgerHeading"))
        for d in report["documents"]:
            story.append(p(f'{d["name"]} · {d["kind"]} · File ID: {d["id"]}'))
    else:
        story.append(p("Document index: No supporting documents for this period. Other unlinked documents remain in the document inbox.", "LedgerSmall"))
    def footer(canvas, document):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#65705c"))
        canvas.drawString(44, 27, "Leah's Ledger · Private financial records")
        canvas.drawRightString(568, 27, f"Page {document.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def word_report(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    style = document.styles["Normal"]
    style.font.name = "Segoe UI"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor.from_string("45543E")
    style.paragraph_format.space_after = Pt(5)
    document.styles["Title"].font.name = "Georgia"
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].font.name = "Georgia"
        document.styles[name].font.color.rgb = RGBColor.from_string("526C4E")
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Heading 2"].font.size = Pt(11)
    for name in ("Title", "Title Char", "Heading 1", "Heading 1 Char", "Heading 2", "Heading 2 Char"):
        if name not in document.styles:
            continue
        selected = document.styles[name]
        selected.font.name = "Georgia"
        # Some bundled templates carry a title rule and theme font overrides.
        for borders in list(selected.element.iter(qn("w:pBdr"))):
            borders.getparent().remove(borders)
        for fonts in selected.element.iter(qn("w:rFonts")):
            for attribute in list(fonts.attrib):
                if attribute.split("}")[-1].lower().endswith("theme"):
                    del fonts.attrib[attribute]
    document.add_heading("Business income and expense report", 0)
    document.add_paragraph(f"{report['owner']} | {report['start']} through {report['end']}")
    document.add_paragraph(f"{len(report['entries'])} transactions | {report['needs_attention']} entries need attention | Generated {datetime.now():%Y-%m-%d %H:%M}")
    document.add_paragraph(REPORT_NOTE)
    for title, rows in report_sections(report):
        document.add_heading(title, 1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Inches(5.5), Inches(1.6)
        for label, value in rows or [("No records in this period", "—")]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
            for cell in cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F1F4E9")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_heading("Transaction details", 1)
    if not report["entries"]:
        document.add_paragraph("No transactions in this period.")
    for e in report["entries"]:
        lines = detail_lines(e, report["documents"])
        document.add_heading(lines[0], 2)
        for line in lines[1:]:
            document.add_paragraph(line)
    if report["documents"]:
        document.add_heading("Document index", 1)
        for d in report["documents"]:
            document.add_paragraph(f'{d["name"]} | {d["kind"]} | File ID: {d["id"]}')
    else:
        document.add_paragraph("Document index: No supporting documents for this period. Other unlinked documents remain in the document inbox.")
    section.footer.paragraphs[0].text = "Leah's Ledger · Private financial records"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class LedgerServer(ThreadingHTTPServer):
    daemon_threads = True
    def __init__(self, address, store):
        self.store = store
        self.token = secrets.token_urlsafe(32)
        self.report_lock = threading.Lock()
        super().__init__(address, Handler)
        self.origin = f"http://127.0.0.1:{self.server_port}"


class Handler(BaseHTTPRequestHandler):
    server_version = "LeahsLedger/1.0"

    def log_message(self, fmt, *args):
        logging.info("%s %s", self.command, urllib.parse.urlparse(self.path).path)

    def start_response(self, length, mime, filename=None, status=200):
        self.send_response(status)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(length))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("X-Frame-Options", "SAMEORIGIN")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' blob: data:; frame-src 'self' blob:; object-src 'none'; base-uri 'none'; frame-ancestors 'self'; form-action 'self'")
        if filename:
            self.send_header("Content-Disposition", "attachment; filename*=UTF-8''" + urllib.parse.quote(filename, safe=""))
        self.end_headers()

    def send_bytes(self, payload, mime, filename=None, status=200):
        self.start_response(len(payload), mime, filename, status)
        self.wfile.write(payload)

    def send_file(self, path, mime, filename):
        self.start_response(path.stat().st_size, mime, filename)
        with path.open("rb") as contents:
            shutil.copyfileobj(contents, self.wfile, 1024 * 1024)

    def json_response(self, value, status=200):
        self.send_bytes(json.dumps(value, ensure_ascii=False).encode(), "application/json; charset=utf-8", status=status)

    def authenticated(self, query):
        token = self.headers.get("X-Ledger-Token") or query.get("token", [""])[0]
        if self.headers.get("Host") != f"127.0.0.1:{self.server.server_port}":
            return False
        origin = self.headers.get("Origin")
        if origin and origin != self.server.origin:
            return False
        return secrets.compare_digest(token, self.server.token)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        query = urllib.parse.parse_qs(parsed.query)
        try:
            if parsed.path in ("/", "/app.js", "/styles.css", "/favicon.svg"):
                if self.headers.get("Host") != f"127.0.0.1:{self.server.server_port}":
                    return self.json_response({"error": "Invalid host."}, 403)
                file = RESOURCES / "ui" / ({"/": "index.html"}.get(parsed.path, parsed.path.lstrip("/")))
                mime = {".html": "text/html; charset=utf-8", ".js": "text/javascript; charset=utf-8", ".css": "text/css; charset=utf-8", ".svg": "image/svg+xml"}[file.suffix]
                return self.send_bytes(file.read_bytes(), mime)
            if not self.authenticated(query):
                return self.json_response({"error": "Open Leah's Ledger from its desktop app to continue."}, 403)
            if parsed.path == "/api/state":
                return self.json_response(self.server.store.snapshot())
            if parsed.path.startswith("/api/documents/"):
                ident = parsed.path.rsplit("/", 1)[-1]
                meta, file = self.server.store.document(ident)
                inline = meta["mime"] in ("image/png", "image/jpeg", "image/gif", "image/webp", "application/pdf", "text/plain") and not query.get("download")
                return self.send_bytes(file.read_bytes(), meta["mime"] if inline else "application/octet-stream", None if inline else meta["name"])
            if parsed.path == "/api/backup":
                with tempfile.TemporaryDirectory(dir=self.server.store.root) as tmp:
                    path = Path(tmp) / "backup.zip"
                    self.server.store.backup(path)
                    return self.send_file(path, "application/zip", f"Leahs-Ledger-backup-{date.today()}.zip")
            if parsed.path == "/api/report":
                report = report_data(self.server.store, query.get("start", [""])[0], query.get("end", [""])[0])
                fmt = query.get("format", ["pdf"])[0]
                filename = f"Leahs-Ledger-{report['start']}-to-{report['end']}"
                with self.server.report_lock:
                    if fmt == "pdf":
                        return self.send_bytes(pdf_report(report), "application/pdf", filename + ".pdf")
                    if fmt == "docx":
                        return self.send_bytes(word_report(report), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename + ".docx")
                    if fmt == "csv":
                        return self.send_bytes(csv_report(report), "text/csv; charset=utf-8", filename + ".csv")
                    if fmt == "zip":
                        with tempfile.TemporaryDirectory(dir=self.server.store.root) as tmp:
                            output = Path(tmp) / "report.zip"
                            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                                archive.writestr("Summary.pdf", pdf_report(report))
                                archive.writestr("Summary.docx", word_report(report))
                                archive.writestr("Transactions.csv", csv_report(report))
                                index = []
                                for d in report["documents"]:
                                    meta, file = self.server.store.document(d["id"])
                                    archive.write(file, "Documents/" + d["id"] + "-" + d["name"])
                                    index.append({"id": d["id"], "original_name": d["name"], "transaction_ids": d["transaction_ids"]})
                                archive.writestr("Document-index.json", json.dumps(index, indent=2, ensure_ascii=False))
                                archive.writestr("Read me.txt", "This package includes transactions in the selected period, their linked original documents, and bank/platform statements and tax forms assigned to the selected year(s). Other unlinked receipts and contracts remain in the document inbox. A complete backup contains all records and files.\n\n" + REPORT_NOTE)
                            return self.send_file(output, "application/zip", filename + ".zip")
                raise ValueError("Choose PDF, Word, CSV, or a complete report package.")
            return self.json_response({"error": "Not found."}, 404)
        except ValueError as error:
            self.json_response({"error": str(error)}, 400)
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            logging.exception("Request failed")
            self.json_response({"error": "Something went wrong. Your saved records are safe. Please try again."}, 500)

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        if not self.authenticated(urllib.parse.parse_qs(parsed.query)):
            return self.json_response({"error": "This request was not authorized by the desktop app."}, 403)
        try:
            length = int(self.headers.get("Content-Length", 0))
            if parsed.path == "/api/restore" and self.headers.get("Content-Type") == "application/zip":
                if not 0 < length <= 2 * 1024 * 1024 * 1024:
                    raise ValueError("Choose a backup ZIP up to 2 GB.")
                with tempfile.TemporaryDirectory(dir=self.server.store.root) as tmp:
                    path = Path(tmp) / "restore.zip"
                    remaining = length
                    with path.open("wb") as output:
                        while remaining:
                            chunk = self.rfile.read(min(remaining, 1024 * 1024))
                            if not chunk:
                                raise ValueError("The backup upload was interrupted. Please choose the file again.")
                            output.write(chunk)
                            remaining -= len(chunk)
                    return self.json_response(self.server.store.restore_archive(path))
            limit = 350 * 1024 * 1024 if parsed.path == "/api/restore" else 40 * 1024 * 1024
            if length <= 0 or length > limit:
                raise ValueError("This file is too large. Try a smaller file.")
            data = json.loads(self.rfile.read(length))
            if not isinstance(data, dict):
                raise ValueError("Invalid request.")
            store = self.server.store
            routes = {"/api/transactions": store.save_transaction, "/api/documents": store.upload,
                      "/api/settings": store.save_settings, "/api/restore": store.restore,
                      "/api/import/preview": store.csv_preview, "/api/import": store.csv_import,
                      "/api/transactions/remove": lambda d: store.remove_transaction(d.get("id")),
                      "/api/transactions/undo": lambda d: store.undo_remove(d.get("id"))}
            if parsed.path == "/api/backup-now":
                store.auto_backup(force=True)
                return self.json_response({"saved": True})
            if parsed.path == "/api/shutdown":
                store.auto_backup(force=True)
                self.json_response({"closed": True})
                threading.Thread(target=self.server.shutdown, daemon=True).start()
                return
            handler = routes.get(parsed.path)
            if handler is None:
                return self.json_response({"error": "Not found."}, 404)
            result = handler(data)
            if parsed.path not in ("/api/import/preview", "/api/restore"):
                store.auto_backup(on_change=True)
            self.json_response(result)
        except (ValueError, json.JSONDecodeError) as error:
            self.json_response({"error": str(error)}, 400)
        except (BrokenPipeError, ConnectionResetError):
            pass
        except Exception:
            logging.exception("Save failed")
            self.json_response({"error": "This change could not be saved. Your earlier records are safe. Please try again."}, 500)


def edge_path():
    for variable in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
        path = Path(os.environ.get(variable, "C:/")) / "Microsoft/Edge/Application/msedge.exe"
        if path.is_file():
            return path
    return None


def launch_window(url, data_folder):
    edge = edge_path()
    if edge:
        return subprocess.Popen([str(edge), "--app=" + url, "--window-size=1380,920", "--no-first-run",
                                 "--user-data-dir=" + str(data_folder / "window-profile")],
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    import webbrowser
    webbrowser.open(url)
    return None


def main():
    parser = argparse.ArgumentParser(description=APP_NAME)
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--no-window", action="store_true")
    parser.add_argument("--port", type=int, default=0)
    options = parser.parse_args()
    data_folder = options.data_dir or Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "LeahsLedger"
    data_folder.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(filename=data_folder / "app.log", level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    # A file lock prevents two desktop instances from editing/restoring one ledger.
    lock_file = open(data_folder / "app.lock", "a+b")
    if os.name == "nt":
        import msvcrt
        try:
            lock_file.seek(0)
            lock_file.write(b"0")
            lock_file.flush()
            lock_file.seek(0)
            msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
        except OSError:
            try:
                active = json.loads((data_folder / "active.json").read_text())
                launch_window(active["url"], data_folder)
            except Exception:
                pass
            return
    store = Store(data_folder)
    server = LedgerServer(("127.0.0.1", options.port), store)
    url = server.origin + "/#token=" + server.token
    (data_folder / "active.json").write_text(json.dumps({"url": url}), encoding="utf-8")
    if options.no_window:
        print(url, flush=True)
    else:
        launch_window(url, data_folder)
    try:
        server.serve_forever(poll_interval=.3)
    except KeyboardInterrupt:
        pass
    finally:
        store.auto_backup(force=True)
        server.server_close()
        (data_folder / "active.json").unlink(missing_ok=True)
        lock_file.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        logging.exception("Could not start")
        if getattr(sys, "frozen", False) or sys.executable.lower().endswith("pythonw.exe"):
            import ctypes
            ctypes.windll.user32.MessageBoxW(0, f"Leah's Ledger could not start.\n\n{error}", APP_NAME, 0x10)
        else:
            raise
